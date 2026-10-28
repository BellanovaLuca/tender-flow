from collections.abc import AsyncIterator
from pathlib import Path
from typing import TypeVar

import docx
import pytest
from pydantic import BaseModel

from app import knowledge_base
from app.schemas import (
    CompanyProfile,
    Criterion,
    DocumentReview,
    SearchPlan,
    SectionMapping,
    SectionPlan,
    SectionReview,
    TenderRequirements,
)
from app.settings import Settings

T = TypeVar("T", bound=BaseModel)


def make_pdf(text: str) -> bytes:
    """PDF minimale a una pagina con testo estraibile (pypdf tollera xref approssimativa)."""
    content = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += b"%d 0 obj\n" % number + obj + b"\nendobj\n"
    xref = len(body)
    body += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    body += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    body += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return body


class FakeLlm:
    """LLM deterministico: la prima validazione di A.1 fallisce per esercitare il ciclo di revisione."""

    def __init__(self, inconsistent: bool = False) -> None:
        self.reviews: dict[str, int] = {}
        self.writer_prompts: list[str] = []
        self.document_reviews = 0
        # Se True, la prima rilettura dell'offerta segnala un'incoerenza tra A.1 e B.
        self.inconsistent = inconsistent

    async def generate_structured(self, *, system: str, prompt: str, schema: type[T]) -> T:
        if schema is TenderRequirements:
            return schema.model_validate(
                TenderRequirements(
                    contracting_authority="Ente Test",
                    tender_object="Servizi RSA",
                    technical_max_points=30,
                    economic_max_points=70,
                    technical_threshold=18,
                    page_limit=5,
                    formatting_rules="Arial 11",
                    mandatory_attachments=["Scheda formazione"],
                    criteria=[
                        Criterion(id="A.1", title="Modello organizzativo", description="turni", max_points=20, scoring_type="D"),
                        Criterion(id="B", title="Formazione", description="piano", max_points=10, scoring_type="D"),
                    ],
                ).model_dump()
            )
        if schema is CompanyProfile:
            return schema.model_validate({"summary": "Cooperativa sociale", "certifications": ["UNI ISO 45001:2018"], "facts": []})
        if schema is SearchPlan:
            return schema.model_validate({"queries": ["modello organizzativo turni"]})
        if schema is SectionPlan:
            return schema.model_validate(
                SectionPlan(mappings=[SectionMapping(criterion_id="A.1", template_section_ids=["T1"], rationale="x")]).model_dump()
            )
        if schema is SectionReview:
            section = "A.1" if "CRITERIO A.1" in prompt else "B"
            self.reviews[section] = self.reviews.get(section, 0) + 1
            failed = section == "A.1" and self.reviews[section] == 1
            review = SectionReview(
                is_compliant=not failed, expected_coefficient=0.8, issues=["manca turno"] if failed else [], strengths=[]
            )
            return schema.model_validate(review.model_dump())
        if schema is DocumentReview:
            self.document_reviews += 1
            if self.inconsistent and self.document_reviews == 1:
                issue = {"section_ids": ["A.1", "B"], "description": "ore del coordinatore diverse", "severity": "bloccante"}
                return schema.model_validate({"is_consistent": False, "issues": [issue]})
            return schema.model_validate(DocumentReview(is_consistent=True, issues=[]).model_dump())
        raise AssertionError(f"schema inatteso {schema}")

    async def stream_text(self, *, system: str, prompt: str) -> AsyncIterator[str]:
        self.writer_prompts.append(prompt)
        pieces = ["Il modello ", "organizzativo ", "prevede turni."]
        if "CRITERIO B " in prompt and "CONFERMATE DALL'OPERATORE" not in prompt:
            pieces.append(" Ore settimanali [DA CONFERMARE: 12 ore].")
            pieces.append("\n- Il **Coordinatore** opera secondo OHSAS 18001.")
        for piece in pieces:
            yield piece


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    kb = tmp_path / "kb"
    offer = kb / "Storico" / "Gara1" / "Offerta_Tecnica"
    offer.mkdir(parents=True)
    (offer / "offerta.pdf").write_bytes(make_pdf("Modello organizzativo con turni e formazione del personale"))
    (kb / "Template").mkdir()
    template = docx.Document()
    template.add_heading("Modello organizzativo", level=1)
    template.add_paragraph("Direttore di struttura e coordinatore infermieristico.")
    template.save(str(kb / "Template" / "TEST_template.docx"))
    sample = kb / "Capitolato target"
    sample.mkdir()
    (sample / "Disciplinare.pdf").write_bytes(make_pdf("Criteri di valutazione offerta tecnica"))
    knowledge_base._kb = None
    return Settings(gemini_api_key="test", knowledge_base_dir=kb, data_dir=tmp_path / "data")  # type: ignore[arg-type]


@pytest.fixture
def anyio_backend() -> str:
    return "asyncio"
