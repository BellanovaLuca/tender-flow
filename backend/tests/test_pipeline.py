import asyncio
from pathlib import Path
from typing import TypeVar

import docx
import pytest

from app.agents.review import _known_section
from app.llm import TokenUsage
from pydantic import BaseModel

from app.pipeline import resumable_sections, resume_section, run_pipeline
from app.runs import Run, RunStore
from app.settings import Settings
from tests.conftest import FakeLlm

T = TypeVar("T", bound=BaseModel)


@pytest.mark.anyio
async def test_pipeline_runs_all_agents_and_revises_failed_section(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))

    llm = FakeLlm()
    await run_pipeline(run, paths, llm, settings)  # type: ignore[arg-type]

    types = [e.type for e in run.events]
    assert types[0] == "run.started" and types[-1] == "run.completed"
    assert "run.failed" not in types
    final_status = {
        e.data["section_id"]: e.data["status"] for e in run.events if e.type == "section.status"
    }
    # B cita sempre una norma superata: la correzione dell'Offer Validator fallisce e passa all'operatore.
    assert final_status == {"A.1": "approved", "B": "flagged"}
    operator = [e.data["status"] for e in run.events if e.type == "agent.status" and e.agent == "operator"]
    assert operator[-1] == "warning"
    resets = [e.data["section_id"] for e in run.events if e.type == "section.reset"]
    assert resets == ["A.1", "B"]
    score = next(e for e in run.events if e.type == "score.estimated")
    assert score.data["total"] == pytest.approx(24.0)
    assert run.document_path is not None and Path(run.document_path).stat().st_size > 0
    assert [e.seq for e in run.events] == list(range(len(run.events)))
    checks = next(e for e in run.events if e.type == "template.validation").data["checks"]
    by_label = {c["label"]: c for c in checks}  # type: ignore[index, union-attr]
    assert by_label["Dati da confermare"]["items"] == [{"section_id": "B", "text": "B: [DA CONFERMARE: 12 ore]"}]
    assert by_label["Norme superate"]["items"] == [
        {"section_id": "B", "text": "B: cita OHSAS 18001 (sostituita da UNI ISO 45001)"}
    ]
    paragraphs = docx.Document(str(run.document_path)).paragraphs
    bullet = next(p for p in paragraphs if "Coordinatore" in p.text)
    assert bullet.text == "Il Coordinatore opera secondo OHSAS 18001."
    assert [r.text for r in bullet.runs if r.bold] == ["Coordinatore"]
    assert llm.writer_prompts and all("Certificazioni attuali: UNI ISO 45001:2018" in p for p in llm.writer_prompts)


@pytest.mark.anyio
async def test_follow_replays_from_cursor(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    await run.publish("run.started")
    await run.publish("log", message="x")
    await run.publish("run.completed")

    replayed = [e.type async for e in run.follow(after=0) if e is not None]

    assert replayed == ["log", "run.completed"]


def test_token_usage_cost_bills_thinking_as_output() -> None:
    usage = TokenUsage(calls=1, input_tokens=1_000_000, output_tokens=100_000, thinking_tokens=100_000)

    assert usage.cost_usd(input_price_per_m=0.30, output_price_per_m=2.50) == pytest.approx(0.80)


@pytest.mark.anyio
async def test_operator_input_resumes_chain_from_writer(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))
    llm = FakeLlm()
    await run_pipeline(run, paths, llm, settings)  # type: ignore[arg-type]
    first_completion = len(run.events)

    await resume_section(run, "B", ["[DA CONFERMARE: 12 ore] → 14 ore settimanali"])

    resumed = run.events[first_completion:]
    types = [e.type for e in resumed]
    assert types[0] == "run.resumed" and types[-1] == "run.completed"
    assert "template.validation" in types and "score.estimated" in types
    assert "14 ore settimanali" in llm.writer_prompts[-1]
    statuses = {e.data["section_id"]: e.data["status"] for e in resumed if e.type == "section.status"}
    assert statuses["B"] == "approved"
    operator = [e.data["status"] for e in resumed if e.type == "agent.status" and e.agent == "operator"]
    assert operator[-1] == "done"
    assert run.finished and run.document_path is not None


@pytest.mark.anyio
async def test_inconsistency_is_fixed_on_lower_weight_section(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))
    llm = FakeLlm(inconsistent=True)

    await run_pipeline(run, paths, llm, settings)  # type: ignore[arg-type]

    coherence_prompts = [p for p in llm.writer_prompts if "Allinea alla sezione A.1" in p]
    assert len(coherence_prompts) == 1 and "CRITERIO B " in coherence_prompts[0]


@pytest.mark.anyio
async def test_minor_inconsistencies_are_not_sent_back(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))
    llm = FakeLlm()

    async def minor_review(*, system: str, prompt: str, schema: type) -> object:
        if schema.__name__ == "DocumentReview":
            issue = {"section_ids": ["A.1", "B"], "description": "certificazioni ripetute", "severity": "minore"}
            return schema.model_validate({"is_consistent": False, "issues": [issue]})
        return await FakeLlm.generate_structured(llm, system=system, prompt=prompt, schema=schema)

    llm.generate_structured = minor_review  # type: ignore[method-assign]
    await run_pipeline(run, paths, llm, settings)  # type: ignore[arg-type]

    assert not [p for p in llm.writer_prompts if "certificazioni ripetute" in p]
    checks = next(e for e in reversed(run.events) if e.type == "template.validation").data["checks"]
    coherence = next(c for c in checks if c["label"] == "Coerenza tra sezioni")  # type: ignore[index]
    assert coherence["status"] == "ok" and coherence["detail"] == "Nessuna incoerenza bloccante · 1 minori"  # type: ignore[index]


def test_invented_subcodes_map_to_existing_criterion() -> None:
    known = {"A.1", "H", "C.2"}

    assert _known_section("H.2", known) == "H"
    assert _known_section("C.2.1", known) == "C.2"
    assert _known_section("A.1", known) == "A.1"
    assert _known_section("Z.9", known) is None


@pytest.mark.anyio
async def test_uploaded_template_replaces_default(settings: Settings, tmp_path: Path) -> None:
    custom = docx.Document()
    custom.add_paragraph("1. Organizzazione del servizio")
    custom.add_paragraph("Testo proveniente dal template caricato per questa gara.")
    template_path = tmp_path / "Template_Offerta_2026.docx"
    custom.save(str(template_path))
    run = RunStore(settings.runs_dir).create()
    run.template_path = template_path
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))
    llm = FakeLlm()

    await run_pipeline(run, paths, llm, settings)  # type: ignore[arg-type]

    a1_prompt = next(p for p in llm.writer_prompts if "CRITERIO A.1" in p)
    assert "Testo proveniente dal template caricato" in a1_prompt
    assert "Direttore di struttura" not in a1_prompt  # testo del template predefinito
    checks = next(e for e in run.events if e.type == "template.validation").data["checks"]
    residual = next(c for c in checks if c["label"] == "Riferimenti a gare precedenti")  # type: ignore[index]
    assert residual["status"] == "ok"  # type: ignore[index]


class CountingLlm(FakeLlm):
    def __init__(self) -> None:
        super().__init__()
        self.structured_calls = 0

    async def generate_structured(self, *, system: str, prompt: str, schema: type[T]) -> T:
        self.structured_calls += 1
        return await super().generate_structured(system=system, prompt=prompt, schema=schema)


async def _wait_for_event(run: Run, type_: str) -> None:
    async def first_match() -> None:
        async for event in run.follow():
            if event is not None and event.type == type_:
                return

    await asyncio.wait_for(first_match(), timeout=10)


@pytest.mark.anyio
async def test_pause_holds_llm_calls_until_continue(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))
    llm = CountingLlm()
    await run.pause()

    run.task = asyncio.create_task(run_pipeline(run, paths, llm, settings))  # type: ignore[arg-type]
    await _wait_for_event(run, "documents.parsed")  # il parser non usa il modello e prosegue

    assert llm.structured_calls == 0
    await run.unpause()
    await run.task

    types = [e.type for e in run.events]
    assert types.index("run.paused") < types.index("run.continued") and types[-1] == "run.completed"
    assert llm.structured_calls > 0


@pytest.mark.anyio
async def test_stop_while_paused_ends_run_without_resume(settings: Settings) -> None:
    run = RunStore(settings.runs_dir).create()
    paths = sorted((settings.knowledge_base_dir / "Capitolato target").glob("*.pdf"))
    llm = CountingLlm()
    await run.pause()
    run.task = asyncio.create_task(run_pipeline(run, paths, llm, settings))  # type: ignore[arg-type]
    await _wait_for_event(run, "documents.parsed")

    run.stop()
    with pytest.raises(asyncio.CancelledError):
        await run.task

    assert run.events[-1].type == "run.stopped" and run.finished
    assert llm.structured_calls == 0
    assert resumable_sections(run.id) is None
    assert run.document_path is None
