from pathlib import Path

import docx

from app.agents.analysis import _requirements_corpus
from app.documents import ParsedDocument, read_template_headings
from app.knowledge_base import _discover


def test_template_without_heading_styles_splits_on_numbered_titles(tmp_path: Path) -> None:
    document = docx.Document()
    for text in ("1. Organizzazione del servizio", "Testo organizzazione.", "1.1 Turni", "Testo turni.", "2. Formazione", "Testo formazione."):
        document.add_paragraph(text)
    path = tmp_path / "template.docx"
    document.save(str(path))

    headings = read_template_headings(path)

    assert [(h.level, h.title) for h in headings] == [
        (1, "1. Organizzazione del servizio"),
        (2, "1.1 Turni"),
        (1, "2. Formazione"),
    ]


def test_template_without_any_structure_becomes_single_section(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Testo libero senza titoli.")
    path = tmp_path / "libero.docx"
    document.save(str(path))

    headings = read_template_headings(path)

    assert len(headings) == 1 and "Testo libero" in headings[0].body


def test_requirements_corpus_ranks_by_content_not_filename() -> None:
    documents = [
        ParsedDocument("allegato_tecnico.pdf", "pdf", 1, "specifiche impianti"),
        ParsedDocument("lettera_invito.pdf", "pdf", 1, "criteri di valutazione: offerta tecnica punti max 70"),
    ]

    corpus, included = _requirements_corpus(documents)

    assert included == ["lettera_invito.pdf", "allegato_tecnico.pdf"]
    assert corpus.index("lettera_invito") < corpus.index("allegato_tecnico")


def test_discover_accepts_any_offer_folder_name(tmp_path: Path) -> None:
    for folder in ("Gara A/Offerta tecnica", "Gara B/OFFERTA_TECNICA_2023", "Gara C/Atti di gara"):
        (tmp_path / "Storico" / folder).mkdir(parents=True)
        (tmp_path / "Storico" / folder / "doc.pdf").write_bytes(b"%PDF-")

    tenders = sorted(tender for tender, _ in _discover(tmp_path))

    assert tenders == ["Gara A", "Gara B"]
