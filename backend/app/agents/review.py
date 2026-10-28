"""Agenti finali: validazione dell'offerta completa, stima del punteggio, assemblaggio DOCX."""

import asyncio
import re
from pathlib import Path

import docx
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from pydantic import JsonValue

from app.agents.context import PipelineContext
from app.schemas import AgentId, CheckItem, DocumentReview

_PLACEHOLDER = re.compile(r"\[DA CONFERMARE[^\]]*\]")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
# Norme ritirate e sostitute: lo storico può contenerle, l'offerta nuova non deve citarle.
_SUPERSEDED_STANDARDS = {
    r"OHSAS\s*18001": "sostituita da UNI ISO 45001",
    r"ISO\s*9001\s*:\s*2008": "sostituita da ISO 9001:2015",
    r"ISO\s*14001\s*:\s*2004": "sostituita da ISO 14001:2015",
}


def _legacy_terms(ctx: PipelineContext) -> list[str]:
    """Nomi delle gare precedenti e del template: non devono comparire nella nuova offerta.
    Un termine presente anche nei documenti del nuovo bando è legittimo (es. 'Offerta' nel nome file)."""
    stem_words = [w for w in re.split(r"[_\W]+", ctx.template().path.stem) if len(w) > 3 and not w.isdigit()]
    tender_text = " ".join(d.text for d in ctx.tender_documents).lower()
    candidates = {*ctx.kb().tenders, *stem_words} - {ctx.settings.company_name}
    return sorted(t for t in candidates if t.lower() not in tender_text)


def _known_section(raw_id: str, known_ids: set[str]) -> str | None:
    """Riporta al criterio esistente i sotto-codici inventati dal modello (es. 'H.2' → 'H')."""
    candidate = raw_id.strip()
    while candidate:
        if candidate in known_ids:
            return candidate
        if "." not in candidate:
            return None
        candidate = candidate.rsplit(".", 1)[0]
    return None


def _check(label: str, items: list[CheckItem], ok_detail: str, warn_detail: str, *, is_ok: bool | None = None) -> dict[str, JsonValue]:
    """Un controllo con riassunto breve e i singoli punti (con la sezione a cui si riferiscono)."""
    passed = not items if is_ok is None else is_ok
    return {
        "label": label,
        "status": "ok" if passed else "warning",
        "detail": ok_detail if passed else warn_detail,
        "items": [item.model_dump() for item in items],
    }


async def validate_document(ctx: PipelineContext) -> dict[str, list[str]]:
    """Controlli sull'offerta completa. Restituisce le correzioni da rimandare al Writer, per sezione."""
    run, agent, requirements = ctx.run, AgentId.TEMPLATE_VALIDATOR, ctx.reqs()
    await run.agent(agent, "running", "Controlli formali e di coerenza")
    checks: list[dict[str, JsonValue]] = []
    fixes: dict[str, list[str]] = {}

    words = sum(len(d.text.split()) for d in ctx.drafts)
    pages = words / ctx.settings.words_per_page
    if requirements.page_limit:
        detail = f"{pages:.1f} facciate stimate su {requirements.page_limit} ammesse ({words} parole)"
        is_within = pages <= requirements.page_limit
        checks.append(_check("Limite facciate", [], detail, detail, is_ok=is_within))
        await run.log(agent, detail, "info" if is_within else "warning")

    flagged = [d for d in ctx.drafts if not (d.review and d.review.is_compliant)]
    approved = f"{len(ctx.drafts) - len(flagged)}/{len(ctx.drafts)} sezioni approvate"
    checks.append(_check(
        "Copertura criteri",
        [
            CheckItem(section_id=d.criterion.id, text=f"{d.criterion.id}: {d.review.issues[0] if d.review and d.review.issues else 'non approvata'}")
            for d in flagged
        ],
        approved,
        f"{approved} · {len(flagged)} da rivedere",
    ))

    residual = [
        CheckItem(section_id=d.criterion.id, text=f"{d.criterion.id}: cita '{term}'")
        for term in _legacy_terms(ctx)
        for d in ctx.drafts
        if re.search(rf"\b{re.escape(term)}\b", d.text, re.IGNORECASE)
    ]
    for item in residual:
        fixes.setdefault(item.section_id or "", []).append(f"Rimuovi il riferimento alla gara precedente ({item.text.split(': ', 1)[1]})")
    checks.append(_check("Riferimenti a gare precedenti", residual, "Nessun riferimento residuo", f"{len(residual)} riferimenti da rimuovere"))
    if residual:
        await run.log(agent, "Riferimenti residui a gare precedenti: " + "; ".join(i.text for i in residual), "warning")

    placeholders = [
        CheckItem(section_id=d.criterion.id, text=f"{d.criterion.id}: {match}")
        for d in ctx.drafts
        for match in _PLACEHOLDER.findall(d.text)
    ]
    checks.append(_check("Dati da confermare", placeholders, "Nessun dato da confermare", f"{len(placeholders)} punti marcati [DA CONFERMARE]"))

    superseded = [
        CheckItem(section_id=d.criterion.id, text=f"{d.criterion.id}: cita {match.group(0)} ({replacement})")
        for pattern, replacement in _SUPERSEDED_STANDARDS.items()
        for d in ctx.drafts
        for match in re.finditer(pattern, d.text, re.IGNORECASE)
    ]
    for item in superseded:
        fixes.setdefault(item.section_id or "", []).append(f"Aggiorna la norma citata: {item.text.split(': ', 1)[1]}")
    checks.append(_check("Norme superate", superseded, "Nessuna norma superata citata", f"{len(superseded)} citazioni da aggiornare"))

    attachments = [CheckItem(section_id=None, text=a) for a in requirements.mandatory_attachments]
    checks.append(_check("Allegati obbligatori", attachments, "Nessun allegato richiesto", f"{len(attachments)} allegati da predisporre"))

    full_text = "\n\n".join(f"{d.criterion.id} {d.criterion.title}\n{d.text}" for d in ctx.drafts)
    review = await ctx.llm.generate_structured(
        system=(
            "Sei il quality manager delle offerte. Rileggi l'offerta tecnica completa e segnala le incoerenze "
            "tra sezioni, ciascuna in una frase con i codici delle sezioni coinvolte. Classifica come 'bloccante' "
            "solo le contraddizioni che una commissione penalizzerebbe: numeri diversi per la stessa cosa (ore, "
            "figure, frequenze), ruoli o responsabilità attribuiti in modo incompatibile, impegni presentati sia "
            "come standard sia come miglioria. Tutto il resto (ripetizioni, sinonimi, differenze di stile) è "
            "'minore'. Massimo 5 issues, le bloccanti per prime."
        ),
        prompt=f"VINCOLI FORMALI: {requirements.formatting_rules}\n\nOFFERTA:\n{full_text}",
        schema=DocumentReview,
    )
    points = {d.criterion.id: d.criterion.max_points for d in ctx.drafts}
    known_ids = set(points)
    for issue in review.issues:
        involved = list(dict.fromkeys(i for i in (_known_section(raw, known_ids) for raw in issue.section_ids) if i))
        # Le minori restano un suggerimento: riscriverle non converge e genererebbe solo nuove differenze.
        if not involved or issue.severity != "bloccante":
            continue
        # La sezione con più punti fa da riferimento; le altre si allineano, così le correzioni non si rincorrono.
        reference = max(involved, key=lambda i: points[i])
        targets = [i for i in involved if i != reference] or [reference]
        for section_id in targets:
            prefix = f"Allinea alla sezione {reference}: " if section_id != reference else ""
            fixes.setdefault(section_id, []).append(prefix + issue.description)
    inconsistencies = [
        CheckItem(
            section_id=next((s for s in (_known_section(i, known_ids) for i in issue.section_ids) if s), None),
            text=(f"[{issue.severity}] {', '.join(issue.section_ids)}: " if issue.section_ids else f"[{issue.severity}] ")
            + issue.description,
        )
        for issue in sorted(review.issues, key=lambda i: i.severity != "bloccante")
    ]
    blocking = sum(i.severity == "bloccante" for i in review.issues)
    checks.append(_check(
        "Coerenza tra sezioni",
        inconsistencies,
        f"Nessuna incoerenza bloccante · {len(inconsistencies)} minori" if inconsistencies else "Nessuna incoerenza rilevata",
        f"{blocking} bloccanti · {len(inconsistencies) - blocking} minori",
        is_ok=blocking == 0,
    ))
    for item in inconsistencies:
        await run.log(agent, item.text, "warning")

    await run.publish("template.validation", agent, checks=checks, words=words, pages=round(pages, 1))
    warnings = sum(c["status"] != "ok" for c in checks)
    await run.agent(agent, "warning" if warnings else "done", f"{len(checks)} controlli · {warnings} avvisi")
    fixes.pop("", None)
    return fixes


async def estimate_score(ctx: PipelineContext) -> None:
    run, agent, requirements = ctx.run, AgentId.SCORER, ctx.reqs()
    await run.agent(agent, "running", "Simulazione commissione")
    rows = [
        {
            "id": d.criterion.id,
            "title": d.criterion.title,
            "max_points": d.criterion.max_points,
            "points": round((d.review.expected_coefficient if d.review else 0) * d.criterion.max_points, 2),
        }
        for d in ctx.drafts
    ]
    total = round(sum(float(r["points"]) for r in rows), 2)
    threshold = requirements.technical_threshold
    await run.publish(
        "score.estimated", agent, rows=rows, total=total, max=requirements.technical_max_points, threshold=threshold
    )
    message = f"Punteggio tecnico stimato {total:g}/{requirements.technical_max_points:g}"
    if threshold:
        message += f" · soglia {threshold:g} {'superata' if total >= threshold else 'NON superata'}"
    await run.log(agent, message, "info" if not threshold or total >= threshold else "warning")
    await run.agent(agent, "done", message)


def _add_runs(paragraph: Paragraph, text: str) -> None:
    """Il modello a volte usa **grassetto** markdown: lo converte in grassetto Word invece di lasciare gli asterischi."""
    for index, part in enumerate(_BOLD.split(text)):
        if part:
            paragraph.add_run(part.replace("*", "")).bold = index % 2 == 1


def _write_docx(ctx: PipelineContext, target: Path) -> None:
    requirements = ctx.reqs()
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.5
    document.add_heading("Offerta tecnica", level=0)
    document.add_paragraph(f"{requirements.tender_object}\n{requirements.contracting_authority}")
    document.add_paragraph(f"Operatore economico: {ctx.settings.company_name}")
    for draft in ctx.drafts:
        document.add_heading(f"{draft.criterion.id} – {draft.criterion.title}", level=1)
        for block in draft.text.split("\n"):
            line = block.strip()
            if not line:
                continue
            is_bullet = re.match(r"^[-*]\s+", line) is not None
            paragraph = document.add_paragraph(style="List Bullet" if is_bullet else None)
            _add_runs(paragraph, re.sub(r"^[-*]\s+", "", line))
    document.save(target)


async def assemble(ctx: PipelineContext) -> Path:
    run, agent = ctx.run, AgentId.ASSEMBLER
    await run.agent(agent, "running", "Composizione documento Word")
    target = run.directory / "Offerta_Tecnica.docx"
    await asyncio.to_thread(_write_docx, ctx, target)
    await run.log(agent, f"Generato {target.name} · {len(ctx.drafts)} sezioni · Arial 11, interlinea 1,5")
    await run.agent(agent, "done", target.name)
    return target
