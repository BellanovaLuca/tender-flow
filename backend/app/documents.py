"""Estrazione testo da PDF/DOCX e ricezione sicura degli upload."""

import asyncio
import hashlib
import json
import logging
import re
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from zipfile import BadZipFile

import docx
import pypdf
from docx.opc.exceptions import PackageNotFoundError
from fastapi import UploadFile
from lxml.etree import XMLSyntaxError

logger = logging.getLogger(__name__)
logging.getLogger("pypdf").setLevel(logging.ERROR)

DocumentKind = Literal["pdf", "docx"]

_MAGIC: dict[bytes, DocumentKind] = {b"%PDF-": "pdf", b"PK\x03\x04": "docx"}
_SAFE_NAME = re.compile(r"[^A-Za-z0-9._ -]")
_CHUNK_BYTES = 1024 * 1024


class UploadRejected(ValueError):
    """Upload non valido; il messaggio è sicuro da mostrare all'utente."""


@dataclass(frozen=True)
class ParsedDocument:
    name: str
    kind: DocumentKind
    pages: int
    text: str


@dataclass(frozen=True)
class Heading:
    id: str
    level: int
    title: str
    body: str


def sanitize_filename(raw: str) -> str:
    name = Path(raw.replace("\\", "/")).name
    name = _SAFE_NAME.sub("_", name).strip(" .")
    return name[:120] or "documento"


def detect_kind(head: bytes) -> DocumentKind | None:
    for magic, kind in _MAGIC.items():
        if head.startswith(magic):
            return kind
    return None


async def save_upload(upload: UploadFile, dest_dir: Path, max_bytes: int) -> Path:
    """Salva in streaming con cap di dimensione; valida il tipo via magic-byte."""
    name = sanitize_filename(upload.filename or "")
    first = await upload.read(_CHUNK_BYTES)
    kind = detect_kind(first)
    if kind is None:
        raise UploadRejected(f"{name}: formato non supportato (ammessi PDF e DOCX)")
    target = dest_dir / f"{Path(name).stem}.{kind}"
    written = 0
    try:
        with target.open("wb") as fh:
            chunk = first
            while chunk:
                written += len(chunk)
                if written > max_bytes:
                    raise UploadRejected(f"{name}: supera la dimensione massima consentita")
                fh.write(chunk)
                chunk = await upload.read(_CHUNK_BYTES)
    except UploadRejected:
        target.unlink(missing_ok=True)
        raise
    return target


async def save_template(upload: UploadFile, dest_dir: Path, max_bytes: int) -> Path:
    """Salva il template aziendale della gara: deve essere un DOCX leggibile."""
    path = await save_upload(upload, dest_dir, max_bytes)
    try:
        if path.suffix != ".docx":
            raise UploadRejected(f"{path.name}: il template deve essere un file Word (.docx)")
        await asyncio.to_thread(docx.Document, str(path))
    except (PackageNotFoundError, BadZipFile, KeyError, XMLSyntaxError) as exc:
        path.unlink(missing_ok=True)
        raise UploadRejected(f"{path.name}: template Word non leggibile") from exc
    except UploadRejected:
        path.unlink(missing_ok=True)
        raise
    return path


def _extract_sync(path: Path) -> ParsedDocument:
    kind: DocumentKind = "docx" if path.suffix.lower() == ".docx" else "pdf"
    if kind == "pdf":
        reader = pypdf.PdfReader(path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return ParsedDocument(path.name, kind, len(reader.pages), text)
    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    return ParsedDocument(path.name, kind, 0, text)


def extract_cached(path: Path, cache_dir: Path) -> ParsedDocument:
    """Estrazione con cache su disco (chiave = hash del contenuto): lo storico va letto una volta sola."""
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    cache_file = cache_dir / f"{digest}.json"
    if cache_file.exists():
        cached = json.loads(cache_file.read_text(encoding="utf-8"))
        return ParsedDocument(path.name, cached["kind"], cached["pages"], cached["text"])
    parsed = _extract_sync(path)
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(
        json.dumps({"kind": parsed.kind, "pages": parsed.pages, "text": parsed.text}), encoding="utf-8"
    )
    return parsed


async def extract(path: Path, cache_dir: Path) -> ParsedDocument:
    return await asyncio.to_thread(extract_cached, path, cache_dir)


_STYLE_HEADING = re.compile(r"(?:Heading|Titolo)\s*(\d)")
# Titoli numerati scritti a mano, es. "1.2. Progetto di gestione" o "A.1 – Modello organizzativo".
_NUMBERED_HEADING = re.compile(r"^(?:\d+(?:\.\d+)*\.?|[A-Z](?:\.\d+)*\s*[–—-]|[A-Z]\d*\.\d+\.?)\s+\S")


def _split_by(paragraphs: list[tuple[str, str]], level_of: Callable[[str, str], int | None]) -> list[Heading]:
    headings: list[Heading] = []
    current: tuple[int, str] | None = None
    body: list[str] = []

    def flush() -> None:
        if current is not None:
            headings.append(Heading(f"T{len(headings) + 1}", current[0], current[1], "\n".join(body)))

    for style, text in paragraphs:
        level = level_of(style, text)
        if level is not None:
            flush()
            current, body = (level, text), []
        elif current is not None:
            body.append(text)
    flush()
    return headings


def _style_level(style: str, _text: str) -> int | None:
    match = _STYLE_HEADING.search(style)
    return int(match.group(1)) if match else None


def _numbered_level(_style: str, text: str) -> int | None:
    if len(text) > 140 or not _NUMBERED_HEADING.match(text):
        return None
    return len([part for part in text.split()[0].split(".") if part])


def read_template_headings(path: Path) -> list[Heading]:
    """Divide il template aziendale in sezioni: stili Titolo di Word, altrimenti titoli numerati,
    altrimenti un'unica sezione con tutto il testo."""
    document = docx.Document(str(path))
    paragraphs = [
        (p.style.name if p.style else "", p.text.strip()) for p in document.paragraphs if p.text.strip()
    ]
    headings = _split_by(paragraphs, _style_level) or _split_by(paragraphs, _numbered_level)
    if headings:
        return headings
    return [Heading("T1", 1, path.stem, "\n".join(text for _, text in paragraphs))]
